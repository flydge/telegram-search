"""Synthetic DOCX safety through real readers and bounded isolated workers.

No provider, runtime or network. Small inert fixtures and child-only resource
probes preserve literal body semantics and prove atomic failure and cleanup.
"""
from __future__ import annotations

import hashlib
import json
import subprocess
import os
import stat
import sys
import tempfile
import unittest
import warnings
import zipfile
from pathlib import Path

from telegram_search_mcp import document_page_reader, document_reader


CT = "http://schemas.openxmlformats.org/package/2006/content-types"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
WORD = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
OFFICE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
REL_MIME = "application/vnd.openxmlformats-package.relationships+xml"
MAIN_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
BODY = '<w:document xmlns:w="' + WORD + '"><w:body><w:p><w:r><w:t>Visible_Aé🙂e\u0301</w:t></w:r></w:p><w:p><w:r><w:t>end</w:t></w:r></w:p></w:body></w:document>'
VISIBLE = "Visible_Aé🙂e\u0301\nend"


class DocxSafetyTests(unittest.TestCase):
    def setUp(self):
        sandbox = tempfile.TemporaryDirectory(prefix='f9c-docx-safety-', dir='/private/tmp')
        self.addCleanup(sandbox.cleanup)
        self.root = Path(sandbox.name)
        self.sequence = 0

    def package(self, *, extras=(), overrides=(), main=BODY, main_rels=None, root_rels=None):
        """Generate a <=2MiB archive; unsafe parts remain tiny or unreachable."""
        content_types = ('<Types xmlns="' + CT + '"><Default Extension="xml" ContentType="application/xml"/>'
                         '<Default Extension="rels" ContentType="' + REL_MIME + '"/>'
                         '<Default Extension="bin" ContentType="application/octet-stream"/>'
                         '<Override PartName="/word/document.xml" ContentType="' + MAIN_MIME + '"/>'
                         + ''.join('<Override PartName="/' + name + '" ContentType="' + mime + '"/>'
                                   for name, mime in overrides) + '</Types>')
        relationships = root_rels or ('<Relationships xmlns="' + REL + '"><Relationship Id="rId1" '
                         'Type="' + OFFICE + 'officeDocument" Target="word/document.xml"/></Relationships>')
        entries = [('[Content_Types].xml', content_types.encode()),
                   ('_rels/.rels', relationships.encode()),
                   ('word/document.xml', main.encode() if isinstance(main, str) else main)]
        if main_rels is not None:
            entries.append(('word/_rels/document.xml.rels', main_rels.encode()))
        entries.extend(extras)
        self.sequence += 1
        path = self.root / f'fixture-{self.sequence}.docx'
        with warnings.catch_warnings():
            warnings.simplefilter('ignore', UserWarning)  # Intentional duplicate-name fixture writing only.
            with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_STORED) as archive:
                for item in entries:
                    name, data = item[:2]
                    compression = item[2] if len(item) == 3 else zipfile.ZIP_STORED
                    archive.writestr(name, data, compress_type=compression)
        self.assertLess(path.stat().st_size, 2 * 1024 * 1024, 'small_fixture_bound')
        return path

    def read(self, path, reader, *, max_chars=20000, offset=0):
        if reader == 'legacy':
            return document_reader.read_document(path, max_chars=max_chars)
        data = path.read_bytes()
        return document_page_reader.read_document_page(
            path, sha256=hashlib.sha256(data).hexdigest(), size_bytes=len(data),
            name=path.name, mime_type=None, max_chars=max_chars, offset=offset, render_pages=False,
        )

    def accepted(self, path, expected=VISIBLE):
        for reader in ('legacy', 'page'):
            with self.subTest(reader=reader):
                result = self.read(path, reader)
                self.assertEqual((result.status, result.text), ('complete', expected))

    def rejected(self, path, *, budget=False, max_chars=20000, offset=0):
        for reader in ('legacy', 'page'):
            with self.subTest(reader=reader):
                try:
                    result = self.read(path, reader, max_chars=max_chars, offset=offset)
                except Exception as error:
                    self.fail(reader + ' raised instead of safe response: ' + type(error).__name__)
                if reader == 'legacy':
                    self.assertEqual(result.status, 'error')
                    self.assertEqual((result.text, result.image_paths), ('', ()))
                else:
                    self.assertIn(result.status, ('limit_reached',) if budget else ('error', 'unsupported'))
                    self.assertEqual((result.text, result.text_start, result.text_end, result.images, result.has_more),
                                     ('', 0, 0, (), False))
                    if budget:
                        self.assertEqual(result.detail, 'docx_expansion_limit')
                    self.assertNotIn('private-marker', result.detail)

    def rels(self, *relations):
        return '<Relationships xmlns="' + REL + '">' + ''.join(relations) + '</Relationships>'

    def relation(self, *, ident='rId2', family='hyperlink', target='https://example.invalid/private-marker', mode='External'):
        return ('<Relationship Id="' + ident + '" Type="' + OFFICE + family + '" Target="' + target
                + '"' + (' TargetMode="' + mode + '"' if mode is not None else '') + '/>')

    def test_plain_unicode_body_and_legacy_prefix_remain_accepted(self):
        path = self.package()
        self.accepted(path)
        result = self.read(path, 'legacy', max_chars=10)
        self.assertEqual((result.status, result.text), ('partial', 'Visible_Aé'))
        page = self.read(path, 'page', max_chars=3, offset=9)
        self.assertEqual((page.status, page.text_start, page.text_end, page.text), ('page', 9, 12, 'é🙂e'))

    def test_ordinary_external_hyperlink_is_inert_and_display_text_survives(self):
        main = ('<w:document xmlns:w="' + WORD + '" xmlns:r="' + OFFICE.rstrip('/') + '"><w:body><w:p>'
                '<w:r><w:t>Visible </w:t></w:r><w:hyperlink r:id="rId2"><w:r><w:t>inert display</w:t></w:r>'
                '</w:hyperlink></w:p></w:body></w:document>')
        path = self.package(main=main, main_rels=self.rels(self.relation()))
        self.accepted(path, 'Visible inert display')

    def test_unknown_benign_internal_family_is_not_a_closed_allowlist(self):
        relations = self.rels('<Relationship Id="rIdBenign" Type="http://example.invalid/relationships/benign" '
                              'Target="media/benign.bin"/>')
        path = self.package(extras=(('word/benign.xml', b'<benign/>'),
                                  ('word/_rels/benign.xml.rels', relations.encode()),
                                  ('word/media/benign.bin', b'benign-bytes')))
        self.accepted(path)

    def test_benign_cycle_and_parent_relative_internal_target_remain_accepted(self):
        relations = self.rels(self.relation(ident='rIdCycle', family='customBenign', target='benign.xml', mode=None),
                              self.relation(ident='rIdParent', family='customBenign', target='../media/benign.bin', mode=None))
        self.accepted(self.package(extras=(('word/benign.xml', b'<benign/>'),
                                          ('word/_rels/benign.xml.rels', relations.encode()),
                                          ('media/benign.bin', b'benign-bytes'))))

    def test_legacy_internal_100000_character_prefix_is_preserved(self):
        for character, length, status in (('z', 100000, 'complete'), ('z', 100001, 'partial'),
                                          ('🙂', 100000, 'complete')):
            with self.subTest(character=character, length=length):
                main = '<w:document xmlns:w="' + WORD + '"><w:body><w:p><w:r><w:t>' + character * length + '</w:t></w:r></w:p></w:body></w:document>'
                result = self.read(self.package(main=main), 'legacy', max_chars=100000)
                self.assertEqual((result.status, len(result.text)), (status, 100000))
                self.assertEqual(result.text, character * 100000)

    def test_safe_stored_body_above_page_ceiling_still_has_legacy_prefix(self):
        main = '<w:document xmlns:w="' + WORD + '"><w:body><w:p><w:r><w:t>' + 'z' * 1000001 + '</w:t></w:r></w:p></w:body></w:document>'
        path = self.package(main=main)
        with zipfile.ZipFile(path) as archive:
            self.assertTrue(all(item.compress_type == zipfile.ZIP_STORED for item in archive.infolist()))
        legacy = self.read(path, 'legacy', max_chars=100000)
        self.assertEqual((legacy.status, len(legacy.text), legacy.text), ('partial', 100000, 'z' * 100000))
        page = self.read(path, 'page', max_chars=1)
        self.assertEqual((page.status, page.text, page.images, page.has_more), ('limit_reached', '', (), False))

    def test_unused_malformed_xml_cannot_publish_harmless_body(self):
        self.rejected(self.package(extras=(('word/unused.xml', b'<private-marker>'),)))

    def test_dtd_is_rejected_in_main_and_unused_xml_in_utf8_and_utf16(self):
        for encoding in ('utf-8', 'utf-16'):
            for location in ('main', 'unused'):
                with self.subTest(encoding=encoding, location=location):
                    xml = (BODY if location == 'main' else '<unused/>')
                    root = 'w:document' if location == 'main' else 'unused'
                    value = '<?xml version="1.0" encoding="' + encoding + '"?><!DOCTYPE ' + root + '>' + xml
                    data = value.encode(encoding)
                    path = self.package(main=data) if location == 'main' else self.package(extras=(('word/unused.xml', data),))
                    self.rejected(path)

    def test_unused_internal_and_external_entity_declarations_are_not_ignored(self):
        declarations = ('<!ENTITY harmless "private-marker">',
                        '<!ENTITY external SYSTEM "https://example.invalid/private-marker">')
        for encoding in ('utf-8', 'utf-16'):
            for declaration in declarations:
                with self.subTest(encoding=encoding, declaration=declaration):
                    # No entity reference, expansion, fetch or large allocation is possible even in baseline.
                    xml = '<?xml version="1.0" encoding="' + encoding + '"?><!DOCTYPE unused [' + declaration + ']><unused/>'
                    self.rejected(self.package(extras=(('word/unused.xml', xml.encode(encoding)),)))

    def test_duplicate_and_case_alias_entries_have_no_ambiguous_interpretation(self):
        with zipfile.ZipFile(self.package()) as archive:
            content_types = archive.read('[Content_Types].xml')
        cases = (('word/document.xml', BODY.encode()), ('word/DOCUMENT.xml', b'<unused/>'),
                 ('_rels/.rels', self.rels('<Relationship Id="rId1" Type="' + OFFICE + 'officeDocument" Target="word/document.xml"/>').encode()),
                 ('[Content_Types].xml', content_types))
        for name, data in cases:
            with self.subTest(name=name):
                self.rejected(self.package(extras=((name, data),)))

    def test_unused_noncanonical_names_cannot_escape_whole_package_audit(self):
        names = ('../escape.xml', '/absolute.xml', 'word/../alias.xml', './relative.xml',
                 'unused//double.xml', 'unused\\windows.xml', 'unused/%2E%2E/alias.xml',
                 'unused/fragment#alias.xml', 'unused/query?alias.xml')
        for name in names:
            with self.subTest(name=name):
                self.rejected(self.package(extras=((name, b'<unused/>'),)))

    def test_unused_special_entries_and_nonempty_directories_are_rejected(self):
        for kind in (stat.S_IFLNK, stat.S_IFCHR, stat.S_IFIFO):
            with self.subTest(mode=kind):
                info = zipfile.ZipInfo('word/special.xml')
                info.create_system = 3
                info.external_attr = (kind | 0o600) << 16
                self.rejected(self.package(extras=((info, b'<unused/>'),)))
        self.rejected(self.package(extras=(('word/nonempty/', b'private-marker'),)))

    def test_unused_unsupported_compression_is_rejected(self):
        for compression in (zipfile.ZIP_BZIP2, zipfile.ZIP_LZMA):
            with self.subTest(compression=compression):
                self.rejected(self.package(extras=(('word/unused.xml', b'<unused/>', compression),)))

    def test_unused_regular_member_crc_is_verified(self):
        path = self.package(extras=(('word/unused.bin', b'private-marker-bytes'),))
        with zipfile.ZipFile(path) as archive:
            info = archive.getinfo('word/unused.bin')
            offset = info.header_offset + 30 + len(info.filename.encode('utf-8')) + len(info.extra)
        with path.open('r+b') as stream:
            stream.seek(offset)
            value = stream.read(1)
            stream.seek(offset)
            stream.write(bytes((value[0] ^ 1,)))
        self.rejected(path)

    def test_small_high_ratio_unused_xml_fails_before_body_publication(self):
        path = self.package(extras=(('word/unused.xml', b'<unused>' + b' ' * 262144 + b'</unused>', zipfile.ZIP_DEFLATED),))
        with zipfile.ZipFile(path) as archive:
            entry = archive.getinfo('word/unused.xml')
            self.assertGreater(entry.file_size, 200 * entry.compress_size)
            self.assertLess(entry.file_size, 300000)
        self.rejected(path, budget=True)

    def test_unused_xml_depth_64_is_accepted_and_65_is_rejected(self):
        shallow = b'<n>' * 64 + b'body' + b'</n>' * 64
        self.accepted(self.package(extras=(('word/unused.xml', shallow),)))
        deep = b'<n>' * 65 + b'body' + b'</n>' * 65
        self.rejected(self.package(extras=(('word/unused.xml', deep),)), budget=True)

    def test_aggregate_nodes_count_across_small_unused_parts(self):
        xml = b'<unused>' + b'<n/>' * 70000 + b'</unused>'
        self.assertLess(len(xml), 300000)
        path = self.package(extras=tuple((f'word/unused{index}.xml', xml) for index in range(3)))
        self.rejected(path, budget=True)

    def test_relationship_content_type_cannot_hide_under_wrong_extension(self):
        for name in ('word/unused.bin', 'word/unused.xml'):
            with self.subTest(name=name):
                self.rejected(self.package(extras=((name, self.rels().encode()),), overrides=((name, REL_MIME),)))

    def test_relationship_namespace_root_nested_and_attributes_are_reserved(self):
        bodies = (self.rels().encode(), ('<unused>' + self.rels() + '</unused>').encode(),
                  ('<unused xmlns:p="' + REL + '" p:Id="private-marker"/>').encode())
        for index, data in enumerate(bodies):
            with self.subTest(case=index):
                self.rejected(self.package(extras=(('word/unused.xml', data),)))

    def test_content_type_namespace_cannot_hide_in_other_part(self):
        bodies = (('<Types xmlns="' + CT + '"/>').encode(),
                  ('<unused><Types xmlns="' + CT + '"/></unused>').encode(),
                  ('<unused xmlns:c="' + CT + '" c:ContentType="private-marker"/>').encode())
        for index, data in enumerate(bodies):
            with self.subTest(case=index):
                self.rejected(self.package(extras=(('word/unused.xml', data),)))

    def test_uppercase_relationship_metadata_alias_is_rejected(self):
        self.rejected(self.package(extras=(('word/_rels/document.xml.RELS', self.rels().encode()),),
                                   overrides=(('word/_rels/document.xml.RELS', REL_MIME),)))

    def test_declared_xml_under_bin_extension_is_audited(self):
        for mime in ('application/xml', 'text/xml', 'application/example+xml'):
            with self.subTest(mime=mime):
                data = b'<!DOCTYPE unused [<!ENTITY harmless "private-marker">]><unused/>'
                self.rejected(self.package(extras=(('word/unused.bin', data),), overrides=(('word/unused.bin', mime),)))

    def test_external_loading_and_active_relation_families_are_rejected(self):
        for family in ('attachedTemplate', 'image', 'oleObject', 'package', 'control', 'vbaProject'):
            with self.subTest(family=family):
                self.rejected(self.package(main_rels=self.rels(self.relation(family=family))))

    def test_unused_active_content_declaration_is_rejected(self):
        for mime in ('application/vnd.ms-office.vbaProject', 'application/vnd.ms-office.activeX+xml',
                     'application/vnd.openxmlformats-officedocument.oleObject'):
            with self.subTest(mime=mime):
                data = b'<unused/>' if mime.endswith('+xml') else b'synthetic-inert-bytes'
                self.rejected(self.package(extras=(('word/unused.bin', data),), overrides=(('word/unused.bin', mime),)))

    def test_unused_relationship_duplicate_ids_and_invalid_modes_are_rejected(self):
        first = self.relation(ident='rId2')
        cases = (self.rels(first, first), self.rels(self.relation(mode='Maybe')),
                 self.rels(self.relation(ident='private-marker-' + 'x' * 1000)))
        for index, relations in enumerate(cases):
            with self.subTest(case=index):
                self.rejected(self.package(extras=(('word/unused.xml', b'<unused/>'),
                                                   ('word/_rels/unused.xml.rels', relations.encode()))))

    def test_unused_internal_targets_require_existing_unambiguous_members(self):
        for target in ('missing.xml', '../../escape.xml', './unused.xml'):
            with self.subTest(target=target):
                relations = self.rels(self.relation(target=target, family='customBenign', mode=None))
                self.rejected(self.package(extras=(('word/unused.xml', b'<unused/>'),
                                                   ('word/_rels/unused.xml.rels', relations.encode()))))

    def test_one_character_prefix_and_nonzero_offset_cannot_bypass_unused_xml_audit(self):
        path = self.package(extras=(('word/unused.xml', b'<!DOCTYPE unused><unused/>'),))
        self.rejected(path, max_chars=1, offset=3)

    def test_existing_500_member_ceiling_is_preserved(self):
        extras = tuple((f'word/unused{index}.xml', b'<unused/>') for index in range(497))
        self.accepted(self.package(extras=extras))
        self.rejected(self.package(extras=extras + (('word/excess.xml', b'<unused/>'),)), budget=True)

    def safety(self):
        import importlib
        try:
            return importlib.import_module('telegram_search_mcp.docx_safety')
        except ImportError:
            self.fail('legacy DOCX isolation is missing')

    def test_actual_legacy_worker_timeout_kills_and_reaps_child(self):
        result = self.injected_worker(self.package(),
                                      'import sys,time;sys.stdin.buffer.read();time.sleep(30)', timeout=.2)
        self.assertEqual((result.status, result.text, result.error), ('error', '', 'worker_timeout'))

    def test_tiny_gridspan_amplification_is_contained_in_actual_legacy_worker(self):
        safety = self.safety()  # Do not run this fixture against the old unbounded legacy implementation.
        main = ('<w:document xmlns:w="' + WORD + '"><w:body><w:tbl><w:tblPr/><w:tblGrid><w:gridCol w:w="1"/>'
                '</w:tblGrid><w:tr><w:tc><w:tcPr><w:gridSpan w:val="1000000000"/></w:tcPr>'
                '<w:p><w:r><w:t>private-marker</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body></w:document>')
        result = safety.read_docx_prefix(self.package(main=main), max_chars=1, timeout=3)
        self.assertEqual((result.status, result.text), ('error', ''))


    def injected_worker(self, path, code, *, max_chars=10, timeout=3, communicate_error=None):
        """Replace only child entrypoint; retain real processes/IPC/parent cleanup."""
        from unittest.mock import patch
        safety = self.safety()
        actual = subprocess.Popen
        children, directories = [], []
        def launch(args, **kwargs):
            output = args[args.index('--legacy-worker') + 1]
            directories.append(Path(output).parent)
            child = actual([sys.executable, '-I', '-B', '-c', code, output], **kwargs)
            children.append(child)
            if communicate_error is not None:
                def interrupted(**options):
                    raise communicate_error()
                child.communicate = interrupted
            return child
        try:
            with patch.object(safety.subprocess, 'Popen', side_effect=launch):
                return safety.read_docx_prefix(path, max_chars=max_chars, timeout=timeout)
        finally:
            self.assertEqual(len(children), 1)
            self.assertIsNotNone(children[0].poll(), 'child remains alive')
            self.assertTrue(children[0].stdin.closed)
            self.assertFalse(directories[0].exists(), 'private output directory leaked')

    def test_legacy_worker_malformed_responses_cannot_publish_text(self):
        path = self.package()
        valid = dict(status='complete', text='safe', processed_bytes=path.stat().st_size,
                     total_bytes=path.stat().st_size, error=None)
        responses = [dict(valid, status='unsupported'), dict(valid, text='x' * 11),
                     dict(valid, text='\ud800'), dict(valid, processed_bytes=True),
                     dict(valid, processed_bytes=0), dict(valid, total_bytes=0),
                     dict(valid, total_bytes=True), dict(valid, error='private-marker'),
                     dict(valid, status='partial'), dict(valid, status='error', error='parser_error'),
                     dict(valid, images=[]), dict(valid, text=12), {}, [],
                     dict(valid, status=['complete']), dict(valid, status='error', text='',
                                                           processed_bytes=0, error=['parser_error'])]
        for index, response in enumerate(responses):
            with self.subTest(case=index):
                code = ('import sys;from pathlib import Path;sys.stdin.buffer.read();'
                        'Path(sys.argv[1]).write_bytes(' + json.dumps(response).encode().__repr__() + ')')
                result = self.injected_worker(path, code)
                self.assertEqual((result.status, result.text, result.processed_bytes), ('error', '', 0))
                self.assertNotIn('private-marker', result.error)

    def test_legacy_worker_truncated_overflow_nonzero_and_duplicate_outputs_fail_atomically(self):
        path = self.package()
        size = path.stat().st_size
        duplicate = ('{"status":"error","status":"complete","text":"safe",'
                     '"processed_bytes":' + str(size) + ',"total_bytes":' + str(size) + ',"error":null}').encode()
        payloads = (b'{"status":', b'x' * (1024 * 1024 + 1), duplicate)
        for index, payload in enumerate(payloads):
            with self.subTest(case=index):
                # Build the overflow in child to keep test command/JSON bounded.
                expression = "b'x'*(1024*1024+1)" if index == 1 else repr(payload)
                code = 'import sys;from pathlib import Path;sys.stdin.buffer.read();Path(sys.argv[1]).write_bytes(' + expression + ')'
                result = self.injected_worker(path, code)
                self.assertEqual((result.status, result.text), ('error', ''))
        valid = dict(status='complete', text='safe', processed_bytes=size, total_bytes=size, error=None)
        code = ('import sys;from pathlib import Path;sys.stdin.buffer.read();Path(sys.argv[1]).write_bytes('
                + repr(json.dumps(valid).encode()) + ');sys.exit(7)')
        self.assertEqual(self.injected_worker(path, code).status, 'error')

    def test_legacy_worker_interruption_and_pipe_failure_kill_reap_and_cleanup(self):
        path = self.package()
        code = 'import sys,time;sys.stdin.buffer.read();time.sleep(30)'
        with self.assertRaises(KeyboardInterrupt):
            self.injected_worker(path, code, communicate_error=KeyboardInterrupt)
        result = self.injected_worker(path, code, communicate_error=BrokenPipeError)
        self.assertEqual((result.status, result.text), ('error', ''))

    def test_legacy_worker_source_symlink_and_after_parse_drift_fail_atomically(self):
        safety = self.safety()
        path = self.package()
        link = self.root / 'link.docx'
        link.symlink_to(path)
        result = safety.read_docx_prefix(link, max_chars=10)
        self.assertEqual((result.status, result.text), ('error', ''))
        code = ('import sys,json;from pathlib import Path;r=json.loads(sys.stdin.buffer.read());'
                'Path(r["path"]).write_bytes(b"changed");'
                'Path(sys.argv[1]).write_text(json.dumps(dict(status="complete",text="safe",'
                'processed_bytes=r["size_bytes"],total_bytes=r["size_bytes"],error=None)))')
        result = self.injected_worker(path, code)
        self.assertEqual((result.status, result.text, result.error), ('error', '', 'artifact_changed'))

    def test_old_timestamp_source_is_valid(self):
        path = self.package()
        os.utime(path, ns=(-1_000_000_000, -1_000_000_000))
        self.accepted(path)

    def test_fixed_legacy_siblings_ignore_hostile_cwd_and_pythonpath(self):
        from unittest.mock import patch
        safety = self.safety()
        path = self.package()
        hostile = self.root / 'hostile'
        hostile.mkdir()
        for name in ('docx.py', 'document_page_reader.py', 'docx_safety.py'):
            (hostile / name).write_text('raise RuntimeError("hostile module loaded")\n')
        previous = Path.cwd()
        try:
            os.chdir(hostile)
            with patch.dict(os.environ, {'PYTHONPATH': str(hostile)}):
                result = safety.read_docx_prefix(path, max_chars=100)
        finally:
            os.chdir(previous)
        self.assertEqual((result.status, result.text), ('complete', VISIBLE))

    def test_actual_legacy_worker_private_request_is_closed_and_typed(self):
        safety = self.safety()
        path = self.package()
        metadata = path.stat()
        valid = dict(version=1, mode='docx_prefix', path=str(path),
                     fingerprint=[metadata.st_dev, metadata.st_ino, metadata.st_mode, metadata.st_uid,
                                  metadata.st_size, metadata.st_mtime_ns, metadata.st_ctime_ns],
                     size_bytes=metadata.st_size, max_chars=10)
        requests = [dict(valid, version=True), dict(valid, mode='page'), dict(valid, offset=0),
                    dict(valid, max_chars=True), dict(valid, max_chars=100001), dict(valid, max_chars=0),
                    dict(valid, size_bytes=True), dict(valid, path='relative.docx'),
                    dict(valid, fingerprint=[True] * 7), dict(valid, fingerprint=[0] * 7)]
        encoded = [json.dumps(value).encode() for value in requests]
        encoded += [b'x' * (16 * 1024 + 1), b'{"version":1,',
                    json.dumps(valid).replace('"version": 1', '"version": 2, "version": 1').encode()]
        for index, raw in enumerate(encoded):
            with self.subTest(case=index):
                output = self.root / ('private-result-' + str(index))
                process = subprocess.run([sys.executable, '-I', '-B', safety.__file__, '--legacy-worker',
                                          str(output), '3'], input=raw, stdout=subprocess.DEVNULL,
                                         stderr=subprocess.DEVNULL, timeout=5)
                self.assertEqual(process.returncode, 0)
                result = json.loads(output.read_bytes())
                self.assertEqual((result['status'], result['text'], result['processed_bytes']), ('error', '', 0))

    def test_actual_legacy_worker_sets_memory_cpu_fd_core_and_output_limits(self):
        safety = self.safety()
        path = self.package()
        # Inject only extraction work in an actual isolated child; the production
        # entry point still sets limits, validates JSON and pins source first.
        loader = ('import importlib.util,sys;from pathlib import Path;'
                  's=importlib.util.spec_from_file_location("probe_safety",' + repr(safety.__file__) + ');'
                  'm=importlib.util.module_from_spec(s);sys.modules[s.name]=m;s.loader.exec_module(m)\n')
        probes = {
            'memory': 'import mmap; mapping=mmap.mmap(-1, 2*1024*1024*1024)',
            'fd': 'handles=[open("/dev/null","rb") for _ in range(100)]',
            'cpu': 'while True: pass',
            'output': 'with open(str(Path(sys.argv[1]).with_name("oversize")),"wb") as out: out.write(b"x"*(2*1024*1024))',
            'core': 'import resource; limits=resource.getrlimit(resource.RLIMIT_CORE);\n if limits==(0,0): raise ValueError("worker_limit")',
        }
        success = ' return m.PrefixResult("complete","uncapped",len(data),len(data))\n'
        baseline = loader + 'def probe(data,max_chars):\n' + success + 'm._prefix=probe\nm._worker_main(sys.argv[1], .3)\n'
        baseline_result = self.injected_worker(path, baseline)
        self.assertEqual((baseline_result.status, baseline_result.text), ('complete', 'uncapped'))
        for name, body in probes.items():
            with self.subTest(limit=name):
                code = loader + 'def probe(data,max_chars):\n ' + body + '\n' + success
                code += 'm._prefix=probe\nm._worker_main(sys.argv[1], .3)\n'
                result = self.injected_worker(path, code, timeout=3)
                self.assertEqual((result.status, result.text), ('error', ''))
                # CPU must be killed by its kernel limit before the parent 3s deadline.
                if name == 'cpu':
                    self.assertEqual(result.error, 'worker_limit')

    def test_real_body_table_merged_cells_empty_skipping_and_lf_are_preserved(self):
        from docx import Document
        document = Document()
        document.add_paragraph('before')
        document.add_paragraph('')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'left'
        table.cell(0, 1).text = 'right'
        table.cell(1, 0).merge(table.cell(1, 1)).text = 'merged'
        document.add_paragraph('after')
        path = self.root / 'real-body.docx'
        document.save(path)
        self.accepted(path, 'before\nleft\nright\nmerged\nmerged\nafter')

    def test_unused_internal_active_relationships_are_not_hidden_by_benign_mime(self):
        for family in ('vbaProject', 'vbaData', 'activeXControlBinary', 'oleObject', 'package',
                       'attachedTemplate', 'control', 'ctrlProp'):
            with self.subTest(family=family):
                relations = self.rels(self.relation(family=family, target='benign.bin', mode=None))
                self.rejected(self.package(extras=(('word/benign.bin', b'inert-bytes'),), main_rels=relations))

    def test_unused_vba_data_content_type_is_active_even_without_relationship(self):
        self.rejected(self.package(extras=(('word/unused.bin', b'<unused/>'),),
                                   overrides=(('word/unused.bin', 'application/vnd.ms-word.vbaData+xml'),)))

    def test_deep_private_json_output_is_atomic_error(self):
        code = ('import sys;from pathlib import Path;sys.stdin.buffer.read();'
                'Path(sys.argv[1]).write_bytes(b"["*10000+b"0"+b"]"*10000)')
        try:
            result = self.injected_worker(self.package(), code)
        except Exception as error:
            self.fail('malformed output escaped parent: ' + type(error).__name__)
        self.assertEqual((result.status, result.text), ('error', ''))

    def test_actual_legacy_child_rechecks_source_after_snapshot_and_parse(self):
        safety = self.safety()
        path = self.package()
        code = ('import importlib.util,sys;'
                's=importlib.util.spec_from_file_location("drift_safety",' + repr(safety.__file__) + ');'
                'm=importlib.util.module_from_spec(s);sys.modules[s.name]=m;s.loader.exec_module(m)\n'
                'original=m._snapshot\n'
                'def snapshot(path,expected):\n'
                ' data,fd=original(path,expected)\n'
                ' path.write_bytes(b"changed-after-snapshot")\n'
                ' return data,fd\n'
                'm._snapshot=snapshot\nm._worker_main(sys.argv[1],3)\n')
        result = self.injected_worker(path, code)
        self.assertEqual((result.status, result.text, result.error), ('error', '', 'artifact_changed'))

    def test_per_xml_byte_budget_uses_stored_small_archive_without_ratio_masking(self):
        path = self.package()
        # Planned postimage <9MiB (below the 10MiB history guardrail). Stream
        # directly into a private stored archive; never build/display a big diff.
        with zipfile.ZipFile(path, 'a', compression=zipfile.ZIP_STORED) as archive:
            with archive.open('word/unused.xml', 'w') as output:
                output.write(b'<unused>')
                for _ in range(128):
                    output.write(b'z' * 65536)
                output.write(b'</unused>')
        self.assertLess(path.stat().st_size, 9 * 1024 * 1024)
        with zipfile.ZipFile(path) as archive:
            self.assertTrue(all(item.compress_type == zipfile.ZIP_STORED for item in archive.infolist()))
        self.rejected(path, budget=True)

    def test_expanded_namespace_attribute_character_budget_has_small_stored_xml(self):
        # 34k expanded 1k namespace attribute names exceed 32MiB of characters,
        # while the serialized XML is <1MiB, below byte/node/depth/ratio limits.
        xml = ('<unused xmlns:x="urn:' + 'q' * 1000 + '" '
               + ' '.join('x:a' + str(index) + '="v"' for index in range(34000)) + '/>').encode()
        self.assertLess(len(xml), 1024 * 1024)
        self.rejected(self.package(extras=(('word/unused.xml', xml),)), budget=True)

    def test_unused_internal_double_leading_slash_target_is_not_canonical(self):
        for target, accepted in (('//word/document.xml', False), ('/word/document.xml', True)):
            with self.subTest(target=target):
                relations = self.rels(self.relation(target=target, family='customBenign', mode=None))
                path = self.package(extras=(('word/unused.xml', b'<unused/>'),
                                           ('word/_rels/unused.xml.rels', relations.encode())))
                self.accepted(path) if accepted else self.rejected(path)

    def test_pathological_vertical_merge_is_contained_in_actual_workers(self):
        self.safety()  # Never exercise pathological library recursion in an old unbounded parent.
        first = '<w:tr><w:tc><w:tcPr><w:vMerge w:val="restart"/></w:tcPr><w:p><w:r><w:t>root</w:t></w:r></w:p></w:tc></w:tr>'
        continuation = '<w:tr><w:tc><w:tcPr><w:vMerge/></w:tcPr><w:p/></w:tc></w:tr>'
        main = ('<w:document xmlns:w="' + WORD + '"><w:body><w:tbl><w:tblPr/>'
                '<w:tblGrid><w:gridCol w:w="1"/></w:tblGrid>' + first + continuation * 1500
                + '</w:tbl></w:body></w:document>')
        path = self.package(main=main)
        legacy = self.safety().read_docx_prefix(path, max_chars=1, timeout=.5)
        self.assertEqual((legacy.status, legacy.text), ('error', ''))
        data = path.read_bytes()
        page = document_page_reader.read_document_page(
            path, sha256=hashlib.sha256(data).hexdigest(), size_bytes=len(data),
            name=path.name, mime_type=None, render_pages=False, max_chars=1, timeout=.5)
        self.assertIn(page.status, ('error', 'limit_reached'))
        self.assertEqual((page.text, page.images, page.has_more), ('', (), False))

    def test_private_json_decoder_recursion_failure_is_atomic_across_python_versions(self):
        from unittest.mock import patch
        safety = self.safety()
        code = ('import sys;from pathlib import Path;sys.stdin.buffer.read();'
                'Path(sys.argv[1]).write_bytes(b"["*10000+b"0"+b"]"*10000)')
        # Python JSON implementations differ in their stack behavior. Keep real
        # child output/IPC/cleanup while reproducing the decoder's documented
        # RecursionError at that dependency boundary on every supported version.
        try:
            with patch.object(safety.json, 'loads', side_effect=RecursionError):
                result = self.injected_worker(self.package(), code)
        except Exception as error:
            self.fail('decoder exception escaped parent: ' + type(error).__name__)
        self.assertEqual((result.status, result.text), ('error', ''))

    def test_selected_docx_broker_mapping_uses_real_store_without_provider_or_runtime(self):
        from telegram_search_mcp.artifact_store import ArtifactStore
        from telegram_search_mcp.broker import Broker
        from telegram_search_mcp.schemas import EvidenceAnchor, ReadAttachmentRequest
        class FakeProvider:
            def __init__(self):
                raise AssertionError('local selected read must not instantiate a provider')
        store = ArtifactStore(cache_dir=self.root / 'broker-artifacts')
        socket_path = self.root / 'never-started.sock'
        broker = Broker(socket_path=socket_path, client_factory=FakeProvider,
                        artifact_store=store, download_source_root=self.root)
        self.addCleanup(broker._executor.shutdown)
        anchor = EvidenceAnchor(chat_id=17, message_id=1024)
        for bad in (False, True):
            with self.subTest(malformed=bad):
                extras = (('word/unused.xml', b'<!DOCTYPE unused><unused/>'),) if bad else ()
                source = self.package(extras=extras)
                artifact = store.store(source)
                broker._artifact_metadata[artifact.artifact_id] = (anchor, 'selected.docx',
                    'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'document')
                read = broker._read_attachment(ReadAttachmentRequest(artifact_id=artifact.artifact_id, max_chars=100))
                self.assertEqual(read.source_anchor, anchor)
                self.assertEqual(read.images, [])
                if bad:
                    self.assertEqual((read.status, read.text, read.processed_bytes, read.coverage_complete),
                                     ('error', '', 0, False))
                else:
                    self.assertEqual((read.status, read.text, read.processed_bytes, read.total_bytes,
                                      read.total_pages, read.coverage_complete),
                                     ('complete', VISIBLE, source.stat().st_size, source.stat().st_size, None, True))
                    prefix = broker._read_attachment(ReadAttachmentRequest(artifact_id=artifact.artifact_id, max_chars=10))
                    self.assertEqual((prefix.status, prefix.text, prefix.processed_bytes, prefix.coverage_complete),
                                     ('partial', 'Visible_Aé', source.stat().st_size, False))
        self.assertFalse(socket_path.exists())

    def test_unused_full_office_package_mimes_are_embedded_packages(self):
        mimes = ('application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                 'application/vnd.openxmlformats-officedocument.wordprocessingml.template',
                 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                 'application/vnd.openxmlformats-officedocument.spreadsheetml.template',
                 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
                 'application/vnd.openxmlformats-officedocument.presentationml.slideshow',
                 'application/vnd.openxmlformats-officedocument.presentationml.template')
        for mime in mimes:
            with self.subTest(mime=mime):
                self.rejected(self.package(extras=(('word/unused.bin', b'inert-inner-package-marker'),),
                                           overrides=(('word/unused.bin', mime),)))

    def test_unknown_benign_part_mime_remains_admissible(self):
        self.accepted(self.package(extras=(('word/unused.bin', b'<benign/>'),),
                                   overrides=(('word/unused.bin', 'application/example-benign+xml'),)))

if __name__ == '__main__':
    unittest.main()
