"""Completed uploads retain the metadata required by real local readers."""
from __future__ import annotations

import base64
import hashlib
import os
import tempfile
import threading
import unittest
from contextlib import contextmanager, nullcontext
from functools import partial
from pathlib import Path
from unittest.mock import patch

from docx import Document
from PIL import Image

from telegram_search_mcp.artifact_store import ArtifactStore
from telegram_search_mcp.broker import Broker
from telegram_search_mcp.broker_client import BrokerClient
from telegram_search_mcp.config import RuntimePolicy
from telegram_search_mcp.local_upload import LocalUploadRegistry
from telegram_search_mcp.schemas import (
    AppendLocalUploadRequest, BeginLocalUploadRequest, FinishLocalUploadRequest,
    ReadAttachmentRequest,
)
from telegram_search_mcp.spreadsheet_models import ReadSpreadsheetRequest
from telegram_search_mcp.presentation_models import ReadPresentationRequest
from xlsx_fixtures import write_xlsx
from pptx_fixtures import make_pptx


class LocalOnlyProvider:
    """Local reader construction may not make any provider requests."""
    def ensure_ready(self):
        pass

    def get_account_id(self):
        return 17

    def request_budget(self, deadline):
        return nullcontext()

    def close(self):
        pass

    def __getattr__(self, name):
        raise AssertionError("unexpected provider operation")


@contextmanager
def running_broker(root: Path, *, policy: RuntimePolicy | None = None):
    policy = policy or RuntimePolicy(enabled_capabilities=("artifacts", "read", "spreadsheets", "presentations"))
    store = ArtifactStore(cache_dir=root / "cache")
    with patch("telegram_search_mcp.broker.LocalUploadRegistry",
               side_effect=partial(LocalUploadRegistry, root=root / "uploads")):
        broker = Broker(socket_path=root / "run" / "broker.sock", artifact_store=store,
                        policy=policy, client_factory=LocalOnlyProvider, verify_peer_uid=False)
    thread = threading.Thread(target=broker.serve_forever)
    thread.start()
    if not broker.wait_until_ready(timeout=2):
        raise AssertionError("broker did not become ready")
    client = BrokerClient(socket_path=root / "run" / "broker.sock", policy=policy,
                          restart_callback=lambda: None, request_timeout=5,
                          retry_backoff_seconds=0)
    try:
        yield broker, client, store
    finally:
        broker.shutdown()
        thread.join(timeout=7)
        if thread.is_alive():
            raise AssertionError("broker did not stop")
        # Process exit closes partial upload descriptors; reproduce that lifecycle
        # in this in-process test utility without adding a production test hook.
        for entry in broker._uploads._entries.values():
            os.close(entry.fd)
        broker._uploads._entries.clear()


def complete_upload(client, source: Path, kind="document"):
    # Small hand-authored fixtures only; helper streaming is tested separately.
    content = source.read_bytes()
    digest = hashlib.sha256(content).hexdigest()
    ready = client.begin_local_upload(BeginLocalUploadRequest(
        file_name=source.name, kind=kind, size_bytes=len(content), sha256=digest))
    if ready.status != "ready":
        raise AssertionError("genuine broker rejected fixture")
    accepted = client.append_local_upload(AppendLocalUploadRequest(
        upload_id=ready.upload_id, index=0, content_base64=base64.b64encode(content).decode("ascii"),
        sha256=digest))
    if accepted.status != "accepted":
        raise AssertionError("genuine broker rejected fixture chunk")
    return client.finish_local_upload(FinishLocalUploadRequest(upload_id=ready.upload_id))


class UploadedArtifactReadTests(unittest.TestCase):
    def test_completed_text_upload_is_readable_with_preserved_name(self):
        with tempfile.TemporaryDirectory(dir="/private/tmp", prefix="f10a-") as directory:
            root = Path(directory)
            source = root / "report.txt"
            source.write_text("Uploaded report evidence", encoding="utf-8")
            with running_broker(root) as (broker, client, _):
                finished = complete_upload(client, source)
                self.assertEqual(finished.status, "complete")
                self.assertEqual(finished.file_name, "report.txt")
                read = client.read_attachment(ReadAttachmentRequest(artifact_id=finished.artifact_id))
                self.assertEqual(read.status, "complete")
                self.assertEqual(read.text, "Uploaded report evidence")
                self.assertIsNone(read.source_anchor)
                self.assertEqual(broker._attachment_metadata(finished.artifact_id),
                                 (None, "report.txt", "text/plain", "document"))

    def test_uploaded_docx_xlsx_pptx_reach_real_supported_readers(self):
        with tempfile.TemporaryDirectory(dir="/private/tmp", prefix="f10a-") as directory:
            root = Path(directory)
            docx = root / "report.docx"
            document = Document()
            document.add_paragraph("Upload DOCX evidence")
            document.save(docx)
            xlsx = write_xlsx(root / "book.xlsx")
            pptx = make_pptx(root / "slides.pptx")
            with running_broker(root) as (broker, client, _):
                word = complete_upload(client, docx)
                self.assertEqual(word.file_name, "report.docx")
                read = client.read_attachment(ReadAttachmentRequest(artifact_id=word.artifact_id))
                self.assertEqual(read.status, "complete")
                self.assertIn("Upload DOCX evidence", read.text)
                workbook = complete_upload(client, xlsx)
                sheet = client.read_spreadsheet(ReadSpreadsheetRequest(
                    artifact_id=workbook.artifact_id, selections=[{"sheet_index": 1, "range": "A1"}]))
                self.assertEqual(sheet.status, "complete")
                self.assertEqual(sheet.cells[0].value, "42")
                self.assertIsNone(sheet.scope.source_anchor)
                deck = complete_upload(client, pptx)
                slides = client.read_presentation(ReadPresentationRequest(artifact_id=deck.artifact_id, slides=[1]))
                self.assertEqual(slides.status, "complete")
                self.assertEqual(slides.slides[0].text, "synthetic slide")
                self.assertIsNone(slides.scope.source_anchor)
                self.assertEqual(broker._attachment_metadata(deck.artifact_id)[1:3],
                    ("slides.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"))

    def test_photo_unknown_document_and_voice_have_honest_metadata(self):
        with tempfile.TemporaryDirectory(dir="/private/tmp", prefix="f10a-") as directory:
            root = Path(directory)
            image = root / "photo.png"
            Image.new("RGB", (4, 4), "blue").save(image)
            unknown = root / "unknown.blob"
            unknown.write_bytes(b"unrecognized document")
            voice = root / "voice.ogg"
            voice.write_bytes(b"synthetic audio")
            with running_broker(root) as (broker, client, _):
                photo = complete_upload(client, image, "photo")
                self.assertEqual(client.read_attachment(ReadAttachmentRequest(artifact_id=photo.artifact_id)).status,
                                 "complete")
                other = complete_upload(client, unknown)
                self.assertEqual(broker._attachment_metadata(other.artifact_id),
                                 (None, "unknown.blob", "application/octet-stream", "document"))
                self.assertEqual(client.read_attachment(ReadAttachmentRequest(artifact_id=other.artifact_id)).status,
                                 "unsupported")
                audio = complete_upload(client, voice, "voice_note")
                self.assertEqual(broker._attachment_metadata(audio.artifact_id)[3], "voice_note")
                self.assertEqual(client.read_attachment(ReadAttachmentRequest(artifact_id=audio.artifact_id)).status,
                                 "unsupported")

    def test_incomplete_finish_and_restarted_registry_never_publish_metadata(self):
        with tempfile.TemporaryDirectory(dir="/private/tmp", prefix="f10a-") as directory:
            root = Path(directory)
            with running_broker(root) as (broker, client, store):
                ready = client.begin_local_upload(BeginLocalUploadRequest(
                    file_name="report.txt", kind="document", size_bytes=3,
                    sha256=hashlib.sha256(b"abc").hexdigest()))
                finished = client.finish_local_upload(FinishLocalUploadRequest(upload_id=ready.upload_id))
                self.assertEqual(finished.status, "error")
                self.assertEqual(broker._artifact_metadata, {})
                self.assertEqual(list(store.cache_dir.glob("artifact_*")), [])
                # Simulate process exit without introducing production cleanup hooks.
                for entry in broker._uploads._entries.values():
                    os.close(entry.fd)
                broker._uploads._entries.clear()
                broker._uploads = LocalUploadRegistry(store, root=root / "uploads")
                rejected = client.append_local_upload(AppendLocalUploadRequest(
                    upload_id=ready.upload_id, index=0, content_base64="YWJj",
                    sha256=hashlib.sha256(b"abc").hexdigest()))
                self.assertEqual(rejected.status, "error")
                self.assertEqual(client.finish_local_upload(FinishLocalUploadRequest(upload_id=ready.upload_id)).status,
                                 "error")
                self.assertEqual(broker._artifact_metadata, {})
