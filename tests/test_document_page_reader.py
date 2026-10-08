from __future__ import annotations

import hashlib
import io
import json
import os
import subprocess
import sys
import tempfile
import time
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import fitz
from docx import Document
from PIL import Image

try:
    from telegram_search_mcp import document_page_reader as reader
except ImportError:
    reader = None


class DocumentPageReaderTests(unittest.TestCase):
    def setUp(self) -> None:
        self.sandbox = tempfile.TemporaryDirectory(
            prefix="parser-tests-", dir=os.environ.get("TELEGRAM_PARSER_TEST_TMP")
        )
        self.addCleanup(self.sandbox.cleanup)
        self.base = Path(self.sandbox.name)

    def _pdf(self, *, width: float = 595, height: float = 842) -> Path:
        path = self.base / "six.pdf"
        with fitz.open() as document:
            for index in range(1, 7):
                page = document.new_page(width=width, height=height)
                page.insert_text((72, 72), f"Page {index}")
            document.save(path)
        return path

    def _read(self, path: Path, **options):
        self.assertIsNotNone(reader, "the approved page reader is missing")
        arguments = {
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "size_bytes": path.stat().st_size,
            "name": path.name,
            "mime_type": None,
            "render_pages": False,
        }
        arguments.update(options)
        return reader.read_document_page(path, **arguments)

    def _reassemble(self, path: Path, *, max_chars: int, **options):
        fragments = []
        offset = 0
        for _ in range(100):
            result = self._read(path, offset=offset, max_chars=max_chars, **options)
            self.assertEqual(result.text_start, offset)
            self.assertEqual(result.text_end, offset + len(result.text))
            fragments.append(result.text)
            if not result.has_more:
                self.assertEqual(result.status, "complete")
                return "".join(fragments)
            self.assertEqual(result.status, "page")
            self.assertGreater(result.text_end, offset)
            offset = result.text_end
        self.fail("text continuation never completed")

    def test_page_six_is_read_without_extracting_the_first_five(self) -> None:
        path = self._pdf()
        if reader is None:
            from telegram_search_mcp.document_reader import read_document
            result = read_document(path, max_pages=5)
        else:
            result = self._read(path, pages=(6,))
        self.assertEqual(result.text, "Page 6\n")
        self.assertEqual(result.status, "complete")
        self.assertEqual(result.selected_pages, (6,))
        self.assertEqual(result.total_pages, 6)

    def test_default_selection_discloses_pages_outside_the_selected_scope(self) -> None:
        result = self._read(self._pdf())
        self.assertEqual(result.selected_pages, (1, 2, 3, 4, 5))
        self.assertEqual(result.total_pages, 6)
        self.assertEqual(result.status, "complete")
        self.assertFalse(result.has_more)
        self.assertEqual(result.text, "Page 1\n\nPage 2\n\nPage 3\n\nPage 4\n\nPage 5\n")

    def test_pdf_selection_order_and_separators_survive_reassembly(self) -> None:
        result = self._reassemble(self._pdf(), pages=(6, 2), max_chars=4)
        self.assertEqual(result, "Page 6\n\nPage 2\n")

    def test_duplicate_empty_and_out_of_range_pdf_selections_fail_atomically(self) -> None:
        path = self._pdf()
        for pages in ((6, 6), (), (1, 7), (0,), (-1,), (True,), (1, 2, 3, 4, 5, 6)):
            with self.subTest(pages=pages):
                result = self._read(path, pages=pages, render_pages=True)
                self.assertEqual(result.status, "invalid_selection")
                self.assertEqual(result.text, "")
                self.assertEqual(result.images, ())
                self.assertFalse(result.has_more)

    def test_previews_use_original_page_numbers_only_at_offset_zero(self) -> None:
        path = self._pdf()
        first = self._read(path, pages=(6, 2), max_chars=3, render_pages=True)
        second = self._read(path, pages=(6, 2), offset=3, max_chars=3, render_pages=True)
        self.assertEqual(tuple(image.page_number for image in first.images), (6, 2))
        self.assertEqual(second.images, ())
        for image in first.images:
            self.assertEqual(image.mime_type, "image/png")
            self.assertLessEqual(len(image.data), 2 * 1024 * 1024)
            with Image.open(io.BytesIO(image.data)) as preview:
                self.assertLessEqual(max(preview.size), 768)

    def test_unicode_offsets_reassemble_combining_nonbmp_and_newline_characters(self) -> None:
        path = self.base / "notes.txt"
        path.write_text("Aé🙂e\u0301\n字Z", encoding="utf-8")
        self.assertEqual(self._reassemble(path, max_chars=3), "Aé🙂e\u0301\n字Z")

    def test_exact_character_boundary_is_complete_and_accepts_end_offset(self) -> None:
        path = self.base / "boundary.txt"
        path.write_text("é🙂", encoding="utf-8")
        result = self._read(path, max_chars=2)
        self.assertEqual((result.status, result.text, result.text_end, result.has_more), ("complete", "é🙂", 2, False))
        end = self._read(path, offset=2)
        self.assertEqual((end.status, end.text_start, end.text_end, end.text), ("complete", 2, 2, ""))
        beyond = self._read(path, offset=3)
        self.assertEqual(beyond.status, "invalid_selection")

    def test_docx_paragraph_and_table_separators_reassemble_exactly(self) -> None:
        path = self.base / "body.docx"
        document = Document()
        document.add_paragraph("Aé🙂")
        document.add_paragraph("")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "left"
        table.cell(0, 1).text = "right"
        document.add_paragraph("end")
        document.save(path)
        self.assertEqual(self._reassemble(path, max_chars=3), "Aé🙂\nleft\nright\nend")

    def test_non_pdf_page_selection_is_rejected(self) -> None:
        path = self.base / "notes.txt"
        path.write_text("selected text")
        result = self._read(path, pages=(1,), render_pages=True)
        self.assertEqual(result.status, "invalid_selection")
        self.assertEqual(result.text, "")
        self.assertEqual(result.images, ())

    def test_image_is_resized_and_validated_without_ocr(self) -> None:
        for format_name, suffix in (("PNG", "png"), ("JPEG", "jpg"), ("WEBP", "webp")):
            with self.subTest(format_name=format_name):
                path = self.base / f"image.{suffix}"
                Image.new("RGB", (1600, 800), "red").save(path, format=format_name)
                result = self._read(path, render_pages=True)
                self.assertEqual(result.status, "complete")
                self.assertEqual(result.kind, "image")
                self.assertEqual(result.text, "")
                self.assertEqual(len(result.images), 1)
                self.assertIsNone(result.images[0].page_number)
                with Image.open(io.BytesIO(result.images[0].data)) as preview:
                    self.assertEqual(preview.size, (768, 384))
                self.assertEqual(self._read(path, render_pages=False).images, ())

    def test_corrupt_image_and_oversized_image_return_no_partial_data(self) -> None:
        path = self.base / "large.png"
        Image.new("1", (5001, 4000)).save(path)
        result = self._read(path, render_pages=True)
        self.assertEqual(result.status, "limit_reached")
        self.assertEqual(result.images, ())
        path.write_bytes(b"private corrupt contents")
        result = self._read(path, render_pages=True)
        self.assertEqual(result.status, "error")
        self.assertEqual(result.images, ())
        self.assertNotIn("private", result.detail)

    def test_password_and_corrupt_pdf_are_safe_errors(self) -> None:
        path = self.base / "secret.pdf"
        with fitz.open() as document:
            document.new_page()
            document.save(path, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="secret", user_pw="secret")
        result = self._read(path, pages=(1,), render_pages=True)
        self.assertEqual((result.status, result.text, result.images), ("error", "", ()))
        self.assertEqual(result.detail, "password_protected")
        path.write_bytes(b"private corrupt contents")
        result = self._read(path, pages=(1,), render_pages=True)
        self.assertEqual((result.status, result.text, result.images), ("error", "", ()))
        self.assertNotIn("private", result.detail)

    def test_giant_pdf_page_is_bounded_without_rendering_partial_data(self) -> None:
        result = self._read(self._pdf(width=2_000_000, height=2_000_000), pages=(1,), render_pages=True)
        self.assertEqual(result.status, "limit_reached")
        self.assertEqual((result.text, result.images), ("", ()))

    def test_hash_size_changes_and_symlinks_are_rejected(self) -> None:
        path = self.base / "notes.txt"
        path.write_text("fixed")
        for options in ({"sha256": "0" * 64}, {"size_bytes": 6}):
            with self.subTest(options=options):
                result = self._read(path, **options)
                self.assertEqual(result.status, "error")
                self.assertEqual((result.text, result.images), ("", ()))
                self.assertEqual(result.detail, "artifact_changed")
        link = self.base / "symlink.txt"
        link.symlink_to(path)
        result = self._read(link)
        self.assertEqual(result.status, "error")
        self.assertEqual(result.text, "")

    def test_extraction_ceiling_returns_no_truncated_scope(self) -> None:
        path = self.base / "too-long.txt"
        path.write_text("🙂" * 1_000_001, encoding="utf-8")
        result = self._read(path, max_chars=2)
        self.assertEqual(result.status, "limit_reached")
        self.assertEqual((result.text, result.images, result.has_more), ("", (), False))
        path.write_text("x" * 1_000_000)
        result = self._read(path, offset=999_999, max_chars=1)
        self.assertEqual((result.status, result.text, result.text_end), ("complete", "x", 1_000_000))

    def test_docx_entry_and_expansion_limits_precede_parsing(self) -> None:
        path = self.base / "bomb.docx"
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for index in range(501):
                archive.writestr(f"item-{index}", b"")
        result = self._read(path)
        self.assertEqual(result.status, "limit_reached")
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            with archive.open("expanded", "w") as output:
                for _ in range(33):
                    output.write(b"x" * (1024 * 1024))
        result = self._read(path)
        self.assertEqual((result.status, result.text, result.images), ("limit_reached", "", ()))

    def test_invalid_utf8_is_safe_and_unknown_kind_is_unsupported(self) -> None:
        path = self.base / "notes.txt"
        path.write_bytes(b"private\xff")
        result = self._read(path)
        self.assertEqual((result.status, result.text, result.images), ("error", "", ()))
        self.assertNotIn("private", result.detail)
        result = self._read(path, name="archive.zip")
        self.assertEqual(result.status, "unsupported")

    def test_input_size_and_invalid_public_arguments_are_bounded(self) -> None:
        path = self.base / "notes.txt"
        path.write_text("data")
        result = self._read(path, size_bytes=64 * 1024 * 1024 + 1)
        self.assertEqual(result.status, "limit_reached")
        for options in ({"max_chars": 0}, {"max_chars": 20_001}, {"offset": -1}, {"max_chars": True}, {"timeout": float("inf")}, {"render_pages": 1}):
            with self.subTest(options=options):
                result = self._read(path, **options)
                self.assertEqual(result.status, "invalid_selection")

    def test_file_change_during_pinned_read_is_rejected(self) -> None:
        self.assertIsNotNone(reader)
        path = self.base / "mutable.txt"
        path.write_text("fixed")
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        original_read = os.read
        changed = False

        def read_then_change(fd, length):
            nonlocal changed
            value = original_read(fd, length)
            if not changed:
                changed = True
                path.write_text("other")
            return value

        with patch.object(reader.os, "read", side_effect=read_then_change):
            with self.assertRaises(reader._ReadFailure) as raised:
                reader._read_pinned(path, digest, 5)
        self.assertEqual(raised.exception.detail, "artifact_changed")

    def test_native_memory_limit_denies_an_allocation_before_parser_import(self) -> None:
        self.assertIsNotNone(reader)
        code = (
            "import runpy\n"
            f"m=runpy.run_path({str(Path(reader.__file__).absolute())!r})\n"
            "fn=m['_set_limits']\n"
            "fn.__globals__['MAX_WORKER_MEMORY_BYTES']=64*1024*1024\n"
            "fn(3.0)\n"
            "try:\n"
            " value=b'x'*(256*1024*1024)\n"
            " print('uncapped')\n"
            "except MemoryError:\n"
            " print('capped')\n"
        )
        result = subprocess.run([sys.executable, "-I", "-B", "-c", code],
                                stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, timeout=5)
        self.assertEqual(result.returncode, 0)
        self.assertEqual(result.stdout, b"capped\n")

    def test_kernel_file_size_limit_caps_worker_output(self) -> None:
        self.assertIsNotNone(reader)
        output_path = self.base / "bounded-output"
        self.assertFalse(output_path.exists())
        code = (
            "import runpy\n"
            f"m=runpy.run_path({str(Path(reader.__file__).absolute())!r})\n"
            "m['_set_limits'](3.0)\n"
            f"with open({str(output_path)!r}, 'wb') as out:\n"
            " for _ in range(3000):\n"
            "  out.write(b'x'*4096)\n"
        )
        result = subprocess.run([sys.executable, "-I", "-B", "-c", code],
                                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=5)
        self.assertNotEqual(result.returncode, 0)
        self.assertGreater(output_path.stat().st_size, 0)
        self.assertLessEqual(output_path.stat().st_size, 9 * 1024 * 1024)

    def test_text_newlines_match_the_existing_utf8_reader(self) -> None:
        path = self.base / "newlines.txt"
        path.write_bytes(b"first\r\nsecond\rthird\n")
        self.assertEqual(self._reassemble(path, max_chars=4), "first\nsecond\nthird\n")

    def test_worker_is_reaped_if_parent_communication_is_interrupted(self) -> None:
        self.assertIsNotNone(reader)
        path = self.base / "notes.txt"
        path.write_text("data")
        script = self.base / "interrupted-worker.py"
        script.write_text("import time\ntime.sleep(10)\n")
        original_popen = subprocess.Popen
        processes = []

        def interrupt_communication(*args, **kwargs):
            process = original_popen(*args, **kwargs)
            processes.append(process)
            process.communicate = lambda **unused: (_ for _ in ()).throw(KeyboardInterrupt())
            return process

        with patch.object(reader, "__file__", str(script)):
            with patch.object(reader.subprocess, "Popen", side_effect=interrupt_communication):
                with self.assertRaises(KeyboardInterrupt):
                    self._read(path)
        process = processes[0]
        reaped = process.poll() is not None
        if not reaped:
            process.kill()
            process.wait()
        if process.stdin is not None:
            process.stdin.close()
        self.assertTrue(reaped, "an interrupted parent left its parser worker running")

    def test_encoded_preview_budgets_discard_all_partial_images(self) -> None:
        path = self.base / "noise.pdf"
        noise = Image.frombytes("RGB", (768, 768), os.urandom(768 * 768 * 3))
        stream = io.BytesIO()
        noise.save(stream, format="PNG")
        with fitz.open() as document:
            for _ in range(4):
                page = document.new_page(width=768, height=768)
                page.insert_image(page.rect, stream=stream.getvalue())
            document.save(path)
        result = self._read(path, render_pages=True)
        self.assertEqual((result.status, result.text, result.images), ("limit_reached", "", ()))
        path = self.base / "noise.png"
        Image.frombytes("RGBA", (768, 768), os.urandom(768 * 768 * 4)).save(path)
        result = self._read(path, render_pages=True)
        self.assertEqual((result.status, result.text, result.images), ("limit_reached", "", ()))

    def test_docx_extraction_ceiling_includes_join_separators(self) -> None:
        path = self.base / "text-limit.docx"
        document = Document()
        for _ in range(4):
            document.add_paragraph("x" * 250_000)
        document.save(path)
        # Stored entries keep the ratio guard from masking the selected-text
        # ceiling: four 250k paragraphs plus three LF separators exceed 1M.
        stored = self.base / "stored-text-limit.docx"
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(stored, "w", compression=zipfile.ZIP_STORED) as target:
            for item in source.infolist():
                target.writestr(item.filename, source.read(item))
        result = self._read(stored)
        self.assertEqual((result.status, result.text, result.images), ("limit_reached", "", ()))
        self.assertEqual(result.detail, "extraction_limit")

    def test_parent_timeout_kills_and_reaps_a_stalled_worker(self) -> None:
        self.assertIsNotNone(reader)
        path = self.base / "notes.txt"
        path.write_text("data")
        script = self.base / "stalled-worker.py"
        script.write_text("import time\ntime.sleep(10)\n")
        started = time.monotonic()
        with patch.object(reader, "__file__", str(script)):
            result = self._read(path, timeout=0.05)
        self.assertEqual(result.status, "limit_reached")
        self.assertEqual(result.detail, "worker_timeout")
        self.assertLess(time.monotonic() - started, 2)
        self.assertEqual((result.text, result.images), ("", ()))


if __name__ == "__main__":
    unittest.main()
