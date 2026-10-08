from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import fitz
from docx import Document
from PIL import Image

from telegram_search_mcp.document_reader import read_document


class DocumentReaderTests(unittest.TestCase):
    def setUp(self) -> None:
        self.sandbox = tempfile.TemporaryDirectory()
        self.addCleanup(self.sandbox.cleanup)
        self.base = Path(self.sandbox.name)

    def test_utf8_text_reports_exact_processed_bytes_and_truncation(self) -> None:
        path = self.base / "notes.py"
        path.write_text("aé🙂z", encoding="utf-8")

        result = read_document(path, max_chars=3)

        self.assertEqual(result.status, "partial")
        self.assertEqual(result.text, "aé🙂")
        self.assertEqual(result.processed_bytes, 7)
        self.assertEqual(result.total_bytes, 8)
        self.assertEqual(result.processed_pages, 0)
        self.assertEqual(result.image_paths, ())

    def test_utf8_text_at_exact_limit_is_complete(self) -> None:
        path = self.base / "notes.txt"
        path.write_text("é🙂", encoding="utf-8")

        result = read_document(path, max_chars=2)

        self.assertEqual(result.status, "complete")
        self.assertEqual(result.text, "é🙂")
        self.assertEqual(result.processed_bytes, 6)

    def test_docx_extracts_paragraphs_and_honors_character_limit(self) -> None:
        path = self.base / "sample.docx"
        document = Document()
        document.add_paragraph("First paragraph")
        document.add_paragraph("Second paragraph")
        document.save(path)

        result = read_document(path, max_chars=20)

        self.assertEqual(result.status, "partial")
        self.assertEqual(result.text, "First paragraph\nSeco")
        self.assertEqual(result.processed_bytes, path.stat().st_size)
        self.assertIsNone(result.total_pages)

    def test_docx_does_not_exceed_limit_when_next_paragraph_requires_separator(self) -> None:
        path = self.base / "boundary.docx"
        document = Document()
        document.add_paragraph("AB")
        document.add_paragraph("C")
        document.save(path)

        result = read_document(path, max_chars=2)

        self.assertEqual(result.status, "partial")
        self.assertEqual(result.text, "AB")

    def test_pdf_extracts_bounded_pages_and_renders_private_page_image(self) -> None:
        path = self.base / "sample.pdf"
        document = fitz.open()
        for value in ("Page one", "Page two"):
            page = document.new_page()
            page.insert_text((72, 72), value)
        document.save(path)
        document.close()

        result = read_document(path, max_pages=1)

        self.assertEqual(result.status, "partial")
        self.assertIn("Page one", result.text)
        self.assertNotIn("Page two", result.text)
        self.assertEqual(result.processed_pages, 1)
        self.assertEqual(result.total_pages, 2)
        self.assertEqual(len(result.image_paths), 1)
        self.assertTrue(result.image_paths[0].exists())
        self.assertNotEqual(result.image_paths[0].parent, path.parent)
        self.assertEqual(result.image_paths[0].stat().st_mode & 0o777, 0o600)
        with Image.open(result.image_paths[0]) as image:
            self.assertEqual(image.format, "PNG")

    def test_image_is_validated_and_returned_without_ocr_claim(self) -> None:
        for format_name, extension in (("JPEG", "jpg"), ("PNG", "png"), ("WEBP", "webp")):
            with self.subTest(format_name=format_name):
                path = self.base / f"pixel.{extension}"
                Image.new("RGB", (2, 2), "red").save(path, format=format_name)

                result = read_document(path)

                self.assertEqual(result.status, "complete")
                self.assertEqual(result.text, "")
                self.assertEqual(result.image_paths, (path,))
                self.assertEqual(result.processed_bytes, path.stat().st_size)

    def test_unknown_format_is_explicitly_unsupported(self) -> None:
        path = self.base / "archive.zip"
        path.write_bytes(b"not a document")

        result = read_document(path)

        self.assertEqual(result.status, "unsupported")
        self.assertEqual(result.processed_bytes, 0)
        self.assertEqual(result.text, "")

    def test_missing_file_is_explicit_error(self) -> None:
        result = read_document(self.base / "missing.txt")

        self.assertEqual(result.status, "error")
        self.assertEqual(result.processed_bytes, 0)
        self.assertTrue(result.error)

if __name__ == "__main__":
    unittest.main()
